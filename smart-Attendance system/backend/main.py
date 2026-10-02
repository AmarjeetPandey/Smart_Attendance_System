"""Pure FastAPI backend for Smart Attendance."""
from contextlib import asynccontextmanager
from pathlib import Path
from datetime import date, datetime
import base64
import csv
import io
import json
import os
import hashlib
import secrets
from collections import defaultdict, deque

from dotenv import load_dotenv
from fastapi import FastAPI, Form, HTTPException, Request, status
from fastapi.responses import HTMLResponse, Response, StreamingResponse
from pydantic import BaseModel
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from sqlalchemy import create_engine, text
from sqlalchemy.exc import IntegrityError
from starlette.middleware.sessions import SessionMiddleware
import qrcode
from itsdangerous import BadSignature, URLSafeTimedSerializer

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")

DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
if not DATABASE_URL:
    raise RuntimeError(
        "DATABASE_URL is required. Copy .env.example to .env and set your online PostgreSQL URL."
    )

if DATABASE_URL.startswith("postgresql://"):
    DATABASE_URL = DATABASE_URL.replace("postgresql://", "postgresql+psycopg://", 1)
if not DATABASE_URL.startswith("postgresql+psycopg://"):
    raise RuntimeError("DATABASE_URL must use PostgreSQL with the postgresql:// or postgresql+psycopg:// scheme.")

engine = create_engine(
    DATABASE_URL,
    pool_pre_ping=True,
    pool_recycle=int(os.getenv("DB_POOL_RECYCLE", "1800")),
    pool_size=int(os.getenv("DB_POOL_SIZE", "5")),
    max_overflow=int(os.getenv("DB_MAX_OVERFLOW", "10")),
    connect_args={"application_name": os.getenv("DB_APPLICATION_NAME", "smart-attendance")},
)
qr_serializer = URLSafeTimedSerializer(os.getenv("SECRET_KEY", "smart-attendance-dev-key"), salt="attendance-qr")
LOGIN_ATTEMPTS: dict[str, deque[float]] = defaultdict(deque)
MAX_LOGIN_ATTEMPTS = 8
LOGIN_WINDOW_SECONDS = 300

SCHEMA = """
CREATE TABLE IF NOT EXISTS admins (
    id SERIAL PRIMARY KEY,
    username TEXT UNIQUE NOT NULL,
    password TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS students (
    id SERIAL PRIMARY KEY,
    name TEXT NOT NULL,
    roll_number TEXT UNIQUE NOT NULL,
    registration_number TEXT UNIQUE NOT NULL,
    course TEXT NOT NULL,
    course_id INTEGER,
    email TEXT DEFAULT '',
    registration_date DATE
);
CREATE TABLE IF NOT EXISTS organizations (
    id SERIAL PRIMARY KEY,
    name TEXT NOT NULL,
    slug TEXT UNIQUE NOT NULL,
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS attendance_policies (
    id SERIAL PRIMARY KEY,
    organization_id INTEGER NOT NULL REFERENCES organizations(id) ON DELETE CASCADE,
    allow_future_dates BOOLEAN NOT NULL DEFAULT FALSE,
    max_backdate_days INTEGER NOT NULL DEFAULT 30,
    qr_valid_seconds INTEGER NOT NULL DEFAULT 86400,
    minimum_percentage NUMERIC(5,2) NOT NULL DEFAULT 75,
    UNIQUE (organization_id)
);
CREATE TABLE IF NOT EXISTS notifications (
    id SERIAL PRIMARY KEY,
    organization_id INTEGER,
    recipient_role TEXT NOT NULL,
    recipient_id INTEGER NOT NULL,
    title TEXT NOT NULL,
    body TEXT NOT NULL,
    read_at TIMESTAMPTZ,
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS attendance (
    id SERIAL PRIMARY KEY,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    attendance_date DATE NOT NULL,
    subject TEXT NOT NULL DEFAULT '',
    status TEXT NOT NULL CHECK (status IN ('Present', 'Absent', 'Late')),
    method TEXT NOT NULL DEFAULT 'Manual',
    CONSTRAINT attendance_student_date_subject_key UNIQUE (student_id, attendance_date, subject)
);
CREATE TABLE IF NOT EXISTS attendance_audit_logs (
    id SERIAL PRIMARY KEY,
    attendance_id INTEGER,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    attendance_date DATE NOT NULL,
    old_status TEXT,
    new_status TEXT NOT NULL,
    method TEXT NOT NULL,
    changed_by_role TEXT NOT NULL,
    changed_by_id INTEGER,
    changed_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS course_catalog (
    id SERIAL PRIMARY KEY,
    name TEXT UNIQUE NOT NULL,
    type TEXT NOT NULL DEFAULT 'Class',
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS leave_requests (
    id SERIAL PRIMARY KEY,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    teacher_id INTEGER,
    date_from DATE NOT NULL,
    date_to DATE NOT NULL,
    reason TEXT NOT NULL,
    status TEXT NOT NULL DEFAULT 'Pending',
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
    reviewed_at TIMESTAMPTZ
);
CREATE TABLE IF NOT EXISTS teachers (
    id SERIAL PRIMARY KEY,
    username TEXT UNIQUE NOT NULL,
    password TEXT NOT NULL,
    name TEXT NOT NULL,
    email TEXT DEFAULT '',
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS teacher_courses (
    teacher_id INTEGER NOT NULL REFERENCES teachers(id) ON DELETE CASCADE,
    course_id INTEGER NOT NULL REFERENCES course_catalog(id) ON DELETE CASCADE,
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (teacher_id, course_id)
);
CREATE TABLE IF NOT EXISTS teacher_class_courses (
    teacher_id INTEGER NOT NULL REFERENCES teachers(id) ON DELETE CASCADE,
    class_id INTEGER NOT NULL REFERENCES course_catalog(id) ON DELETE CASCADE,
    course_id INTEGER NOT NULL REFERENCES course_catalog(id) ON DELETE CASCADE,
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP,
    PRIMARY KEY (teacher_id, class_id, course_id)
);
CREATE TABLE IF NOT EXISTS student_accounts (
    id SERIAL PRIMARY KEY,
    student_id INTEGER UNIQUE NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    username TEXT UNIQUE NOT NULL,
    password TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS exams (
    id SERIAL PRIMARY KEY,
    course TEXT NOT NULL,
    course_id INTEGER,
    subject TEXT NOT NULL,
    exam_date DATE NOT NULL,
    created_by INTEGER REFERENCES admins(id) ON DELETE SET NULL
);
CREATE TABLE IF NOT EXISTS results (
    id SERIAL PRIMARY KEY,
    student_id INTEGER NOT NULL REFERENCES students(id) ON DELETE CASCADE,
    subject TEXT NOT NULL,
    exam_name TEXT NOT NULL,
    marks NUMERIC(6,2) NOT NULL CHECK (marks >= 0),
    total_marks NUMERIC(6,2) NOT NULL CHECK (total_marks > 0),
    grade TEXT NOT NULL DEFAULT '',
    UNIQUE (student_id, subject, exam_name)
);
CREATE TABLE IF NOT EXISTS student_faces (
    student_id INTEGER PRIMARY KEY REFERENCES students(id) ON DELETE CASCADE,
    descriptor JSONB NOT NULL,
    face_url TEXT DEFAULT '',
    updated_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS qr_sessions (
    id SERIAL PRIMARY KEY,
    token_id TEXT UNIQUE NOT NULL,
    attendance_date DATE NOT NULL,
    course TEXT DEFAULT '',
    course_id INTEGER,
    subject TEXT NOT NULL DEFAULT '',
    used BOOLEAN NOT NULL DEFAULT FALSE,
    used_by_student_id INTEGER REFERENCES students(id) ON DELETE SET NULL,
    used_at TIMESTAMPTZ,
    created_at TIMESTAMPTZ NOT NULL DEFAULT CURRENT_TIMESTAMP
);
ALTER TABLE student_faces ADD COLUMN IF NOT EXISTS face_url TEXT DEFAULT '';
ALTER TABLE students ADD COLUMN IF NOT EXISTS course_id INTEGER REFERENCES course_catalog(id) ON DELETE SET NULL;
ALTER TABLE attendance ADD COLUMN IF NOT EXISTS subject TEXT NOT NULL DEFAULT '';
ALTER TABLE exams ADD COLUMN IF NOT EXISTS course_id INTEGER;
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS course_id INTEGER;
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS subject TEXT NOT NULL DEFAULT '';
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS subject TEXT NOT NULL DEFAULT '';
ALTER TABLE leave_requests ADD COLUMN IF NOT EXISTS teacher_id INTEGER REFERENCES teachers(id) ON DELETE SET NULL;
ALTER TABLE leave_requests ALTER COLUMN created_at SET DEFAULT CURRENT_TIMESTAMP;
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS course TEXT DEFAULT '';
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS used BOOLEAN NOT NULL DEFAULT FALSE;
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS used_by_student_id INTEGER REFERENCES students(id) ON DELETE SET NULL;
ALTER TABLE qr_sessions ADD COLUMN IF NOT EXISTS used_at TIMESTAMPTZ;
ALTER TABLE qr_sessions ALTER COLUMN created_at SET DEFAULT CURRENT_TIMESTAMP;
"""


def hash_password(password: str, salt: bytes | None = None) -> str:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, 310_000)
    return f"pbkdf2:sha256:310000${salt.hex()}${digest.hex()}"


def verify_password(password: str, encoded: str) -> bool:
    try:
        algorithm, salt_hex, digest_hex = encoded.split("$", 2)
        algorithm_name, hash_name, iteration_text = algorithm.split(":", 2)
        if algorithm_name != "pbkdf2" or hash_name != "sha256":
            return False
        candidate = hashlib.pbkdf2_hmac(
            "sha256", password.encode(), bytes.fromhex(salt_hex), int(iteration_text)
        )
        return secrets.compare_digest(candidate.hex(), digest_hex)
    except (ValueError, TypeError):
        return False


def require_session(request: Request) -> None:
    login_at = request.session.get("login_at", 0)
    if not request.session.get("account_id") or login_at < datetime.now().timestamp() - 1800:
        request.session.clear()
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Sign in required")
    request.session["login_at"] = datetime.now().timestamp()


def require_role(request: Request, *roles: str) -> None:
    require_session(request)
    if request.session.get("role") not in roles:
        raise HTTPException(status_code=403, detail="This role cannot perform that action")


def parse_date(value: str, field_name: str) -> date:
    try:
        return datetime.strptime(value, "%Y-%m-%d").date()
    except (TypeError, ValueError) as error:
        raise HTTPException(status_code=400, detail=f"{field_name} must be a valid date") from error


def check_login_rate_limit(request: Request) -> None:
    key = request.client.host if request.client else "unknown"
    now = datetime.now().timestamp()
    attempts = LOGIN_ATTEMPTS[key]
    while attempts and attempts[0] <= now - LOGIN_WINDOW_SECONDS:
        attempts.popleft()
    if len(attempts) >= MAX_LOGIN_ATTEMPTS:
        raise HTTPException(status_code=429, detail="Too many login attempts. Try again later.")


def record_login_failure(request: Request) -> None:
    key = request.client.host if request.client else "unknown"
    LOGIN_ATTEMPTS[key].append(datetime.now().timestamp())


def require_teacher_course(request: Request, connection, course: str, subject: str | None = None) -> None:
    if request.session.get("role") != "teacher":
        return
    subject_filter = " AND assigned_course.name = :subject" if subject is not None else ""
    params = {"teacher_id": request.session["account_id"], "course": course}
    if subject is not None:
        params["subject"] = subject
    assigned = connection.execute(text(f"SELECT 1 FROM teacher_class_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.class_id JOIN course_catalog assigned_course ON assigned_course.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND assigned_class.name = :course{subject_filter} LIMIT 1"), params).first()
    if not assigned and subject is None:
        assigned = connection.execute(text("SELECT 1 FROM teacher_courses tc JOIN course_catalog c ON c.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND c.type = 'Class' AND c.name = :course"), {"teacher_id": request.session["account_id"], "course": course}).first()
    if not assigned:
        raise HTTPException(status_code=403, detail="This teacher is not assigned to the selected class and course")


def teacher_student_scope(student_alias: str = "s", subject_column: str | None = None) -> str:
    subject_filter = f" AND assigned_course.name = {subject_column}" if subject_column else ""
    paired = f"EXISTS (SELECT 1 FROM teacher_class_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.class_id JOIN course_catalog assigned_course ON assigned_course.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND (assigned_class.id = {student_alias}.course_id OR assigned_class.name = {student_alias}.course){subject_filter})"
    if subject_column:
        return paired
    legacy = f"EXISTS (SELECT 1 FROM teacher_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND assigned_class.type = 'Class' AND (assigned_class.id = {student_alias}.course_id OR assigned_class.name = {student_alias}.course))"
    return f"({paired} OR {legacy})"


def attendance_policy(connection):
    row = connection.execute(text("SELECT p.* FROM attendance_policies p JOIN organizations o ON o.id = p.organization_id ORDER BY o.id LIMIT 1")).mappings().first()
    return row or {"allow_future_dates": False, "max_backdate_days": 30, "qr_valid_seconds": 86400, "minimum_percentage": 75}


def ensure_integrity_constraints(connection) -> None:
    constraints = (
        ("students", "students_course_id_fkey", "FOREIGN KEY (course_id) REFERENCES course_catalog(id) ON DELETE SET NULL"),
        ("leave_requests", "leave_requests_teacher_id_fkey", "FOREIGN KEY (teacher_id) REFERENCES teachers(id) ON DELETE SET NULL"),
        ("attendance_audit_logs", "attendance_audit_logs_attendance_id_fkey", "FOREIGN KEY (attendance_id) REFERENCES attendance(id) ON DELETE SET NULL"),
        ("notifications", "notifications_organization_id_fkey", "FOREIGN KEY (organization_id) REFERENCES organizations(id) ON DELETE SET NULL"),
        ("exams", "exams_course_id_fkey", "FOREIGN KEY (course_id) REFERENCES course_catalog(id) ON DELETE SET NULL"),
        ("qr_sessions", "qr_sessions_course_id_fkey", "FOREIGN KEY (course_id) REFERENCES course_catalog(id) ON DELETE SET NULL"),
    )
    for table_name, constraint_name, definition in constraints:
        connection.execute(text("""
            DO $$
            BEGIN
                IF NOT EXISTS (
                    SELECT 1 FROM pg_constraint WHERE conname = '%s'
                ) THEN
                    EXECUTE 'ALTER TABLE %s ADD CONSTRAINT %s %s';
                END IF;
            END $$
        """ % (constraint_name, table_name, constraint_name, definition)))
    connection.execute(text("ALTER TABLE attendance DROP CONSTRAINT IF EXISTS attendance_student_id_attendance_date_key"))
    subject_constraint = connection.execute(text("SELECT 1 FROM pg_constraint WHERE conname = 'attendance_student_date_subject_key' AND conrelid = 'attendance'::regclass")).first()
    if not subject_constraint:
        connection.execute(text("ALTER TABLE attendance ADD CONSTRAINT attendance_student_date_subject_key UNIQUE (student_id, attendance_date, subject)"))


def validate_attendance_policy(connection, attendance_date: date) -> None:
    policy = attendance_policy(connection)
    today = datetime.now().date()
    if not policy["allow_future_dates"] and attendance_date > today:
        raise HTTPException(status_code=400, detail="Future attendance dates are disabled")
    if attendance_date < today.fromordinal(today.toordinal() - int(policy["max_backdate_days"] or 30)):
        raise HTTPException(status_code=400, detail="Attendance date is outside the allowed correction window")


def grade_for(marks: float, total_marks: float) -> str:
    percentage = marks / total_marks * 100
    return "A+" if percentage >= 90 else "A" if percentage >= 80 else "B" if percentage >= 70 else "C" if percentage >= 60 else "D" if percentage >= 50 else "F"


@asynccontextmanager
async def lifespan(_: FastAPI):
    admin_username = os.getenv("ADMIN_USERNAME", "").strip()
    admin_password = os.getenv("ADMIN_PASSWORD", "")
    if not admin_username or not admin_password:
        raise RuntimeError("ADMIN_USERNAME and ADMIN_PASSWORD must be set in .env")

    with engine.connect() as connection:
        connection.execute(text("SELECT 1"))
        for statement in SCHEMA.split(";"):
            if statement.strip():
                connection.execute(text(statement))
        ensure_integrity_constraints(connection)
        organization = connection.execute(text("INSERT INTO organizations (name, slug) VALUES ('Default Organization', 'default') ON CONFLICT (slug) DO UPDATE SET name = organizations.name RETURNING id")).mappings().first()
        connection.execute(text("INSERT INTO attendance_policies (organization_id) VALUES (:organization_id) ON CONFLICT (organization_id) DO NOTHING"), {"organization_id": organization["id"]})
        connection.execute(text("UPDATE students s SET course_id = c.id FROM course_catalog c WHERE s.course_id IS NULL AND s.course = c.name"))
        connection.execute(text("UPDATE exams e SET course_id = c.id FROM course_catalog c WHERE e.course_id IS NULL AND e.course = c.name"))
        connection.execute(text("UPDATE qr_sessions q SET course_id = c.id FROM course_catalog c WHERE q.course_id IS NULL AND q.course = c.name"))
        connection.execute(
            text("INSERT INTO admins (username, password) VALUES (:username, :password) ON CONFLICT (username) DO NOTHING"),
            {"username": admin_username, "password": hash_password(admin_password)},
        )
        connection.commit()
    yield
    engine.dispose()


api = FastAPI(title="Smart Attendance API", lifespan=lifespan)
api.mount("/static", StaticFiles(directory=ROOT / "static"), name="static")
api.add_middleware(SessionMiddleware, secret_key=os.getenv("SECRET_KEY", secrets.token_urlsafe(32)), max_age=1800, same_site="lax", https_only=os.getenv("SESSION_HTTPS_ONLY", "0") == "1")
api.add_middleware(
    CORSMiddleware,
    allow_origins=[os.getenv("FRONTEND_ORIGIN", "http://localhost:5173")],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@api.get("/api/health")
def health_check():
    with engine.connect() as connection:
        connection.execute(text("SELECT 1"))
    return {"status": "ok", "database": "postgresql"}


@api.post("/api/login")
def login(request: Request, username: str = Form(...), password: str = Form(...), role: str = Form("admin")):
    check_login_rate_limit(request)
    if role not in {"admin", "teacher", "student"}:
        raise HTTPException(status_code=400, detail="Invalid account role")
    with engine.connect() as connection:
        account = connection.execute(
            text("SELECT id, username, password FROM admins WHERE username = :username"),
            {"username": username.strip()},
        ).mappings().first()
        account_role = "admin"
        account_id = account["id"] if account else None
        if not account:
            account = connection.execute(text("SELECT id, username, password, name FROM teachers WHERE username = :username"), {"username": username.strip()}).mappings().first()
            account_role = "teacher"
            account_id = account["id"] if account else None
        if not account:
            account = connection.execute(text("SELECT sa.id, sa.username, sa.password, sa.student_id, s.name FROM student_accounts sa JOIN students s ON s.id = sa.student_id WHERE sa.username = :username"), {"username": username.strip()}).mappings().first()
            account_role = "student"
            account_id = account["student_id"] if account else None
    if not account or not verify_password(password, account["password"]) or role != account_role:
        record_login_failure(request)
        raise HTTPException(status_code=401, detail="Invalid username or password")
    key = request.client.host if request.client else "unknown"
    LOGIN_ATTEMPTS.pop(key, None)
    display_name = account["name"] if "name" in account and account["name"] else account["username"]
    request.session.update({"account_id": account_id, "admin_id": account_id if account_role == "admin" else None, "username": account["username"], "name": display_name, "role": account_role, "login_at": datetime.now().timestamp()})
    return {"id": account_id, "username": account["username"], "name": display_name, "role": account_role}


@api.post("/api/logout")
def logout(request: Request):
    request.session.clear()
    return {"status": "ok"}


@api.get("/api/students")
def list_students(request: Request, class_id: int | None = None):
    require_session(request)
    with engine.connect() as connection:
        role = request.session["role"]
        if role == "teacher":
            if class_id is None:
                raise HTTPException(status_code=400, detail="Select a class")
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=404, detail="Class not found")
            require_teacher_course(request, connection, class_row["name"])
            rows = connection.execute(text("SELECT * FROM students WHERE course_id = :class_id OR (course_id IS NULL AND course = :class_name) ORDER BY name"), {"class_id": class_id, "class_name": class_row["name"]}).mappings().all()
        elif role == "student":
            rows = connection.execute(text("SELECT * FROM students WHERE id = :student_id ORDER BY name"), {"student_id": request.session["account_id"]}).mappings().all()
        else:
            rows = connection.execute(text("SELECT * FROM students ORDER BY name")).mappings().all()
    return {"students": [dict(row) for row in rows]}


@api.get("/api/courses")
def list_courses(request: Request):
    require_session(request)
    with engine.connect() as connection:
        rows = connection.execute(text("SELECT name, type FROM course_catalog ORDER BY name")).mappings().all()
    return {"courses": [dict(row) for row in rows]}


@api.get("/api/session")
def current_session(request: Request):
    require_session(request)
    return {"id": request.session["account_id"], "username": request.session["username"], "role": request.session["role"]}


@api.get("/api/admin/overview")
def admin_overview(request: Request):
    require_role(request, "admin")
    with engine.connect() as connection:
        teachers = connection.execute(text("SELECT id, username, name, email FROM teachers ORDER BY name")).mappings().all()
        students = connection.execute(text("SELECT s.id, sa.username, s.name, s.roll_number, s.registration_number, s.course, s.course_id, s.email FROM students s LEFT JOIN student_accounts sa ON sa.student_id = s.id ORDER BY s.name")).mappings().all()
        courses = connection.execute(text("SELECT id, name, type FROM course_catalog ORDER BY name")).mappings().all()
        assignments = connection.execute(text("SELECT tc.teacher_id, tc.class_id, tc.course_id, t.name teacher_name, assigned_class.name class_name, assigned_course.name course_name, assigned_course.type course_type, 'pair' assignment_kind FROM teacher_class_courses tc JOIN teachers t ON t.id = tc.teacher_id JOIN course_catalog assigned_class ON assigned_class.id = tc.class_id JOIN course_catalog assigned_course ON assigned_course.id = tc.course_id UNION ALL SELECT tc.teacher_id, NULL::INTEGER class_id, tc.course_id, t.name teacher_name, NULL::TEXT class_name, c.name course_name, c.type course_type, 'legacy' assignment_kind FROM teacher_courses tc JOIN teachers t ON t.id = tc.teacher_id JOIN course_catalog c ON c.id = tc.course_id ORDER BY teacher_name, class_name NULLS LAST, course_name")).mappings().all()
        exams = connection.execute(text("SELECT e.id, c.name course, e.subject, e.exam_date FROM exams e LEFT JOIN course_catalog c ON c.id = e.course_id ORDER BY e.exam_date DESC")).mappings().all()
        counts = connection.execute(text("SELECT (SELECT COUNT(*) FROM students) students, (SELECT COUNT(*) FROM teachers) teachers, (SELECT COUNT(*) FROM course_catalog) courses, (SELECT COUNT(*) FROM exams) exams")).mappings().first()
    return {"counts": dict(counts), "teachers": [dict(row) for row in teachers], "students": [dict(row) for row in students], "courses": [dict(row) for row in courses], "exams": [dict(row) for row in exams], "assignments": [dict(row) for row in assignments]}


@api.post("/api/admin/reset-data")
def reset_admin_data(request: Request, scope: str = Form(...), confirmation: str = Form(...)):
    require_role(request, "admin")
    if scope not in {"admin", "teachers", "students", "ALLCLEAR"}:
        raise HTTPException(status_code=400, detail="Select a valid data group")
    if confirmation != "CLEAR":
        raise HTTPException(status_code=400, detail='Type CLEAR to confirm this destructive action')

    with engine.begin() as connection:
        if scope == "admin":
            connection.execute(text("DELETE FROM exams"))
            connection.execute(text("DELETE FROM course_catalog"))
            connection.execute(text("UPDATE attendance_policies SET allow_future_dates = FALSE, max_backdate_days = 30, qr_valid_seconds = 86400, minimum_percentage = 75"))
            connection.execute(text("DELETE FROM notifications WHERE recipient_role = 'admin'"))
            connection.execute(text("DELETE FROM admins WHERE id <> :admin_id"), {"admin_id": request.session["account_id"]})
        elif scope == "teachers":
            connection.execute(text("DELETE FROM attendance_audit_logs WHERE changed_by_role = 'teacher'"))
            connection.execute(text("DELETE FROM notifications WHERE recipient_role = 'teacher'"))
            connection.execute(text("DELETE FROM teacher_class_courses"))
            connection.execute(text("DELETE FROM teacher_courses"))
            connection.execute(text("UPDATE leave_requests SET teacher_id = NULL WHERE teacher_id IS NOT NULL"))
        elif scope == "students":
            connection.execute(text("DELETE FROM attendance_audit_logs"))
            connection.execute(text("DELETE FROM attendance"))
            connection.execute(text("DELETE FROM results"))
            connection.execute(text("DELETE FROM student_faces"))
            connection.execute(text("DELETE FROM leave_requests"))
            connection.execute(text("DELETE FROM qr_sessions WHERE used_by_student_id IS NOT NULL"))
            connection.execute(text("DELETE FROM notifications WHERE recipient_role = 'student'"))
        else:
            connection.execute(text("TRUNCATE TABLE attendance_audit_logs, attendance, leave_requests, results, student_faces, student_accounts, qr_sessions, notifications, teacher_class_courses, teacher_courses, exams, students, teachers, course_catalog, attendance_policies, organizations RESTART IDENTITY CASCADE"))
            organization = connection.execute(text("INSERT INTO organizations (name, slug) VALUES ('Default Organization', 'default') RETURNING id")).mappings().first()
            connection.execute(text("INSERT INTO attendance_policies (organization_id) VALUES (:organization_id)"), {"organization_id": organization["id"]})

    return {"status": "cleared", "scope": scope}


@api.post("/api/admin/teachers")
def add_teacher(request: Request, username: str = Form(...), password: str = Form(...), name: str = Form(...), email: str = Form("")):
    require_role(request, "admin")
    if len(password) < 8 or not username.strip() or not name.strip():
        raise HTTPException(status_code=400, detail="Teacher name, ID, and an 8-character password are required")
    with engine.begin() as connection:
        row = connection.execute(text("INSERT INTO teachers (username, password, name, email) VALUES (:username, :password, :name, :email) RETURNING id, username, name, email"), {"username": username.strip(), "password": hash_password(password), "name": name.strip(), "email": email.strip()}).mappings().first()
    return dict(row)


@api.put("/api/admin/teachers/{teacher_id}")
def update_teacher(request: Request, teacher_id: int, username: str = Form(...), name: str = Form(...), email: str = Form(""), password: str = Form("")):
    require_role(request, "admin")
    username = username.strip()
    name = name.strip()
    if not username or not name:
        raise HTTPException(status_code=400, detail="Teacher name and ID are required")
    with engine.begin() as connection:
        if password.strip():
            row = connection.execute(text("UPDATE teachers SET username = :username, name = :name, email = :email, password = :password WHERE id = :id RETURNING id, username, name, email"), {"id": teacher_id, "username": username, "name": name, "email": email.strip(), "password": hash_password(password)}).mappings().first()
        else:
            row = connection.execute(text("UPDATE teachers SET username = :username, name = :name, email = :email WHERE id = :id RETURNING id, username, name, email"), {"id": teacher_id, "username": username, "name": name, "email": email.strip()}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Teacher account not found")
    return dict(row)


@api.delete("/api/admin/teachers/{teacher_id}")
def delete_teacher(request: Request, teacher_id: int):
    require_role(request, "admin")
    with engine.begin() as connection:
        result = connection.execute(text("DELETE FROM teachers WHERE id = :id"), {"id": teacher_id})
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Teacher account not found")
    return {"status": "deleted"}


@api.post("/api/admin/courses")
def add_course(request: Request, name: str = Form(...), course_type: str = Form("Class")):
    require_role(request, "admin")
    name = name.strip()
    course_type = course_type.strip() or "Class"
    if not name or course_type not in {"Class", "Subject"}:
        raise HTTPException(status_code=400, detail="A valid class or subject is required")
    with engine.begin() as connection:
        row = connection.execute(text("INSERT INTO course_catalog (name, type) VALUES (:name, :type) ON CONFLICT (name) DO UPDATE SET type = EXCLUDED.type RETURNING id, name, type"), {"name": name, "type": course_type}).mappings().first()
    return dict(row)


@api.get("/api/admin/attendance-policy")
def get_attendance_policy(request: Request):
    require_role(request, "admin")
    with engine.connect() as connection:
        return dict(attendance_policy(connection))


@api.put("/api/admin/attendance-policy")
def update_attendance_policy(request: Request, allow_future_dates: bool = Form(False), max_backdate_days: int = Form(30), qr_valid_seconds: int = Form(86400), minimum_percentage: float = Form(75)):
    require_role(request, "admin")
    if max_backdate_days < 0 or qr_valid_seconds < 60 or not 0 <= minimum_percentage <= 100:
        raise HTTPException(status_code=400, detail="Invalid attendance policy values")
    with engine.begin() as connection:
        organization = connection.execute(text("SELECT id FROM organizations ORDER BY id LIMIT 1")).mappings().first()
        row = connection.execute(text("UPDATE attendance_policies SET allow_future_dates=:future, max_backdate_days=:backdate, qr_valid_seconds=:qr, minimum_percentage=:minimum WHERE organization_id=:organization_id RETURNING *"), {"future": allow_future_dates, "backdate": max_backdate_days, "qr": qr_valid_seconds, "minimum": minimum_percentage, "organization_id": organization["id"]}).mappings().first()
    return dict(row)


@api.post("/api/admin/teacher-courses")
def assign_teacher_course(request: Request, teacher_id: int = Form(...), class_id: int = Form(...), course_id: int = Form(...)):
    require_role(request, "admin")
    with engine.begin() as connection:
        if not connection.execute(text("SELECT 1 FROM teachers WHERE id = :id"), {"id": teacher_id}).first():
            raise HTTPException(status_code=404, detail="Teacher not found")
        if not connection.execute(text("SELECT 1 FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).first():
            raise HTTPException(status_code=400, detail="Select an existing class")
        if not connection.execute(text("SELECT 1 FROM course_catalog WHERE id = :id AND type = 'Subject'"), {"id": course_id}).first():
            raise HTTPException(status_code=400, detail="Select an existing course")
        row = connection.execute(text("INSERT INTO teacher_class_courses (teacher_id, class_id, course_id) VALUES (:teacher_id, :class_id, :course_id) ON CONFLICT DO NOTHING RETURNING teacher_id, class_id, course_id"), {"teacher_id": teacher_id, "class_id": class_id, "course_id": course_id}).mappings().first()
    return dict(row) if row else {"teacher_id": teacher_id, "class_id": class_id, "course_id": course_id}


@api.delete("/api/admin/teacher-courses")
def unassign_teacher_course(request: Request, teacher_id: int, course_id: int, class_id: int | None = None):
    require_role(request, "admin")
    with engine.begin() as connection:
        if class_id is None:
            result = connection.execute(text("DELETE FROM teacher_courses WHERE teacher_id = :teacher_id AND course_id = :course_id"), {"teacher_id": teacher_id, "course_id": course_id})
        else:
            result = connection.execute(text("DELETE FROM teacher_class_courses WHERE teacher_id = :teacher_id AND class_id = :class_id AND course_id = :course_id"), {"teacher_id": teacher_id, "class_id": class_id, "course_id": course_id})
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Assignment not found")
    return {"status": "unassigned"}


@api.put("/api/admin/courses/{course_id}")
def update_course(request: Request, course_id: int, name: str = Form(...), course_type: str = Form("Class")):
    require_role(request, "admin")
    name = name.strip()
    course_type = course_type.strip() or "Class"
    if not name or course_type not in {"Class", "Subject"}:
        raise HTTPException(status_code=400, detail="A valid class or subject is required")
    with engine.begin() as connection:
        row = connection.execute(text("UPDATE course_catalog SET name = :name, type = :type WHERE id = :id RETURNING id, name, type"), {"id": course_id, "name": name, "type": course_type}).mappings().first()
        connection.execute(text("UPDATE students SET course = :name WHERE course_id = :id"), {"id": course_id, "name": name})
    if not row:
        raise HTTPException(status_code=404, detail="Course not found")
    return dict(row)


@api.delete("/api/admin/courses/{course_id}")
def delete_course(request: Request, course_id: int):
    require_role(request, "admin")
    with engine.begin() as connection:
        if connection.execute(text("SELECT 1 FROM students WHERE course_id = :id LIMIT 1"), {"id": course_id}).first():
            raise HTTPException(status_code=409, detail="Move students to another course before deleting this course")
        result = connection.execute(text("DELETE FROM course_catalog WHERE id = :id"), {"id": course_id})
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Course not found")
    return {"status": "deleted"}


@api.post("/api/admin/exams")
def add_exam(request: Request, course: str = Form(...), subject: str = Form(...), exam_date: str = Form(...)):
    require_role(request, "admin")
    course = course.strip()
    subject = subject.strip()
    if not course or not subject or not exam_date:
        raise HTTPException(status_code=400, detail="Course, subject, and exam date are required")
    parsed_date = parse_date(exam_date, "Exam date")
    with engine.begin() as connection:
        course_row = connection.execute(text("SELECT id FROM course_catalog WHERE name = :course"), {"course": course}).mappings().first()
        if not course_row:
            raise HTTPException(status_code=400, detail="Select an existing course")
        row = connection.execute(text("INSERT INTO exams (course, course_id, subject, exam_date, created_by) VALUES (:course, :course_id, :subject, :exam_date, :created_by) RETURNING id, course, subject, exam_date"), {"course": course, "course_id": course_row["id"], "subject": subject, "exam_date": parsed_date, "created_by": request.session["account_id"]}).mappings().first()
    return dict(row)


@api.put("/api/admin/exams/{exam_id}")
def update_exam(request: Request, exam_id: int, course: str = Form(...), subject: str = Form(...), exam_date: str = Form(...)):
    require_role(request, "admin")
    course = course.strip()
    subject = subject.strip()
    if not course or not subject or not exam_date:
        raise HTTPException(status_code=400, detail="Course, subject, and exam date are required")
    parsed_date = parse_date(exam_date, "Exam date")
    with engine.begin() as connection:
        course_row = connection.execute(text("SELECT id FROM course_catalog WHERE name = :course"), {"course": course}).mappings().first()
        if not course_row:
            raise HTTPException(status_code=400, detail="Select an existing course")
        row = connection.execute(text("UPDATE exams SET course = :course, course_id = :course_id, subject = :subject, exam_date = :exam_date WHERE id = :id RETURNING id, course, subject, exam_date"), {"id": exam_id, "course": course, "course_id": course_row["id"], "subject": subject, "exam_date": parsed_date}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Exam schedule not found")
    return dict(row)


@api.delete("/api/admin/exams/{exam_id}")
def delete_exam(request: Request, exam_id: int):
    require_role(request, "admin")
    with engine.begin() as connection:
        result = connection.execute(text("DELETE FROM exams WHERE id = :id"), {"id": exam_id})
    if result.rowcount == 0:
        raise HTTPException(status_code=404, detail="Exam schedule not found")
    return {"status": "deleted"}


@api.get("/api/teacher/overview")
def teacher_overview(request: Request, class_id: int | None = None):
    require_role(request, "teacher", "admin")
    with engine.connect() as connection:
        is_admin = request.session["role"] == "admin"
        if is_admin:
            courses = connection.execute(text("SELECT name, type FROM course_catalog ORDER BY name")).mappings().all()
            classes = connection.execute(text("SELECT id, name FROM course_catalog WHERE type = 'Class' ORDER BY name")).mappings().all()
            class_assignments = []
            scope_params = {}
            student_filter = ""
            result_filter = ""
            leave_filter = ""
        else:
            teacher_id = request.session["account_id"]
            teacher_params = {"teacher_id": teacher_id}
            courses = connection.execute(text("SELECT DISTINCT assigned_class.name, assigned_class.type FROM teacher_class_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.class_id WHERE tc.teacher_id = :teacher_id UNION SELECT DISTINCT assigned_class.name, assigned_class.type FROM teacher_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND assigned_class.type = 'Class' ORDER BY name"), teacher_params).mappings().all()
            classes = connection.execute(text("SELECT DISTINCT assigned_class.id, assigned_class.name FROM teacher_class_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.class_id WHERE tc.teacher_id = :teacher_id UNION SELECT DISTINCT assigned_class.id, assigned_class.name FROM teacher_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND assigned_class.type = 'Class' ORDER BY name"), teacher_params).mappings().all()
            class_assignments = connection.execute(text("SELECT tc.class_id, assigned_class.name class_name, tc.course_id, assigned_course.name course_name, FALSE legacy FROM teacher_class_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.class_id JOIN course_catalog assigned_course ON assigned_course.id = tc.course_id WHERE tc.teacher_id = :teacher_id UNION ALL SELECT tc.course_id class_id, assigned_class.name class_name, NULL::INTEGER course_id, NULL::TEXT course_name, TRUE legacy FROM teacher_courses tc JOIN course_catalog assigned_class ON assigned_class.id = tc.course_id WHERE tc.teacher_id = :teacher_id AND assigned_class.type = 'Class' ORDER BY class_name, course_name"), teacher_params).mappings().all()
            if class_id is None:
                scope_params = {"teacher_id": teacher_id}
                student_filter = " WHERE FALSE"
                result_filter = " WHERE FALSE"
                leave_filter = " AND FALSE"
            else:
                class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
                if not class_row:
                    raise HTTPException(status_code=404, detail="Class not found")
                require_teacher_course(request, connection, class_row["name"])
                scope_params = {"teacher_id": teacher_id, "class_id": class_id, "class_name": class_row["name"]}
                selected_student = "(s.course_id = :class_id OR (s.course_id IS NULL AND s.course = :class_name))"
                student_filter = f" WHERE {teacher_student_scope()} AND {selected_student}"
                result_filter = f" WHERE {teacher_student_scope(subject_column='r.subject')} AND {selected_student}"
                leave_filter = f" AND {teacher_student_scope()} AND {selected_student}"
        students = connection.execute(text(f"SELECT s.id, s.course_id, s.name, s.roll_number, s.registration_number, s.course, s.email, f.face_url, (f.student_id IS NOT NULL) face_enrolled FROM students s LEFT JOIN student_faces f ON f.student_id = s.id{student_filter} ORDER BY s.name"), scope_params).mappings().all()
        results = connection.execute(text(f"SELECT r.id, r.student_id, s.course_id, s.course, s.name student_name, r.subject, r.exam_name, r.marks, r.total_marks, r.grade FROM results r JOIN students s ON s.id = r.student_id{result_filter} ORDER BY r.id DESC LIMIT 100"), scope_params).mappings().all()
        leave_params = {"teacher_id": request.session["account_id"]}
        leave_params.update({key: value for key, value in scope_params.items() if key != "teacher_id"})
        leave_requests = connection.execute(text(f"SELECT lr.id, lr.student_id, s.name student_name, s.roll_number, lr.date_from, lr.date_to, lr.reason, lr.status, lr.created_at FROM leave_requests lr JOIN students s ON s.id = lr.student_id WHERE lr.teacher_id = :teacher_id{leave_filter} ORDER BY lr.created_at DESC"), leave_params).mappings().all()
        attendance = connection.execute(text(f"SELECT a.status, COUNT(*) AS total FROM attendance a JOIN students s ON s.id = a.student_id{student_filter} GROUP BY a.status"), scope_params).mappings().all()
    attendance_summary = {row["status"].lower(): row["total"] for row in attendance}
    return {"students": [dict(row) for row in students], "courses": [dict(row) for row in courses], "classes": [dict(row) for row in classes], "class_assignments": [dict(row) for row in class_assignments], "results": [dict(row) for row in results], "leave_requests": [dict(row) for row in leave_requests], "attendance_summary": attendance_summary}


@api.post("/api/admin/students")
def add_student(request: Request, name: str = Form(...), username: str = Form(...), password: str = Form(...), roll_number: str = Form(...), registration_number: str = Form(...), class_id: int = Form(...), email: str = Form("")):
    require_role(request, "admin")
    name = name.strip()
    username = username.strip()
    roll_number = roll_number.strip()
    registration_number = registration_number.strip()
    if len(password) < 8 or not name or not username or not roll_number or not registration_number:
        raise HTTPException(status_code=400, detail="Student name, ID, class, and an 8-character password are required")
    try:
        with engine.begin() as connection:
            duplicates = connection.execute(text("SELECT EXISTS (SELECT 1 FROM student_accounts WHERE username = :username) AS username_exists, EXISTS (SELECT 1 FROM students WHERE roll_number = :roll) AS roll_exists, EXISTS (SELECT 1 FROM students WHERE registration_number = :registration) AS registration_exists"), {"username": username, "roll": roll_number, "registration": registration_number}).mappings().first()
            if duplicates["username_exists"]:
                raise HTTPException(status_code=409, detail="Student ID is already in use")
            if duplicates["roll_exists"]:
                raise HTTPException(status_code=409, detail="Roll number is already in use")
            if duplicates["registration_exists"]:
                raise HTTPException(status_code=409, detail="Registration number is already in use")
            course_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :class_id AND type = 'Class'"), {"class_id": class_id}).mappings().first()
            if not course_row:
                raise HTTPException(status_code=400, detail="Select a class created by the admin")
            student = connection.execute(text("INSERT INTO students (name, roll_number, registration_number, course, course_id, email, registration_date) VALUES (:name, :roll, :registration, :course, :course_id, :email, CURRENT_DATE) RETURNING id, name, roll_number, course"), {"name": name, "roll": roll_number, "registration": registration_number, "course": course_row["name"], "course_id": course_row["id"], "email": email.strip()}).mappings().first()
            connection.execute(text("INSERT INTO student_accounts (student_id, username, password) VALUES (:student_id, :username, :password)"), {"student_id": student["id"], "username": username, "password": hash_password(password)})
    except IntegrityError as error:
        raise HTTPException(status_code=409, detail="Student ID, roll number, or registration number is already in use") from error
    return dict(student)


@api.put("/api/admin/students/{student_id}")
def update_admin_student(request: Request, student_id: int, name: str = Form(...), username: str = Form(...), roll_number: str = Form(...), registration_number: str = Form(...), class_id: int = Form(...), email: str = Form(""), password: str = Form("")):
    require_role(request, "admin")
    name, username = name.strip(), username.strip()
    roll_number, registration_number = roll_number.strip(), registration_number.strip()
    password = password.strip()
    if not name or not username or not roll_number or not registration_number:
        raise HTTPException(status_code=400, detail="Student name, ID, roll number, and registration number are required")
    if password and len(password) < 8:
        raise HTTPException(status_code=400, detail="A new password must contain at least 8 characters")

    try:
        with engine.begin() as connection:
            student = connection.execute(text("SELECT id FROM students WHERE id = :id FOR UPDATE"), {"id": student_id}).first()
            if not student:
                raise HTTPException(status_code=404, detail="Student not found")
            duplicates = connection.execute(text("SELECT EXISTS (SELECT 1 FROM student_accounts WHERE username = :username AND student_id <> :id) AS username_exists, EXISTS (SELECT 1 FROM students WHERE roll_number = :roll AND id <> :id) AS roll_exists, EXISTS (SELECT 1 FROM students WHERE registration_number = :registration AND id <> :id) AS registration_exists"), {"id": student_id, "username": username, "roll": roll_number, "registration": registration_number}).mappings().first()
            if duplicates["username_exists"] or duplicates["roll_exists"] or duplicates["registration_exists"]:
                raise HTTPException(status_code=409, detail="Student ID, roll number, or registration number is already in use")
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :class_id AND type = 'Class'"), {"class_id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=400, detail="Select a valid class")
            account = connection.execute(text("SELECT student_id FROM student_accounts WHERE student_id = :id"), {"id": student_id}).first()
            if not account and not password:
                raise HTTPException(status_code=400, detail="Set a password to create this student's login")
            row = connection.execute(text("UPDATE students SET name = :name, roll_number = :roll, registration_number = :registration, course = :course, course_id = :class_id, email = :email WHERE id = :id RETURNING id, name, roll_number, registration_number, course, course_id, email"), {"id": student_id, "name": name, "roll": roll_number, "registration": registration_number, "course": class_row["name"], "class_id": class_row["id"], "email": email.strip()}).mappings().first()
            if account and password:
                connection.execute(text("UPDATE student_accounts SET username = :username, password = :password WHERE student_id = :id"), {"id": student_id, "username": username, "password": hash_password(password)})
            elif account:
                connection.execute(text("UPDATE student_accounts SET username = :username WHERE student_id = :id"), {"id": student_id, "username": username})
            else:
                connection.execute(text("INSERT INTO student_accounts (student_id, username, password) VALUES (:id, :username, :password)"), {"id": student_id, "username": username, "password": hash_password(password)})
    except IntegrityError as error:
        raise HTTPException(status_code=409, detail="Student ID, roll number, or registration number is already in use") from error
    return {**dict(row), "username": username}


@api.delete("/api/admin/students/{student_id}")
def delete_admin_student(request: Request, student_id: int):
    require_role(request, "admin")
    with engine.begin() as connection:
        row = connection.execute(text("DELETE FROM students WHERE id = :id RETURNING id, name"), {"id": student_id}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Student not found")
    return {"deleted": True, **dict(row)}


@api.post("/api/teacher/attendance")
def mark_attendance(request: Request, student_id: int = Form(...), class_id: int = Form(...), subject: str = Form(...), attendance_date: str = Form(...), status_value: str = Form(...), method: str = Form("Manual")):
    require_role(request, "teacher", "admin")
    parsed_date = parse_date(attendance_date, "Attendance date")
    if status_value not in {"Present", "Absent", "Late"}:
        raise HTTPException(status_code=400, detail="Invalid attendance status")
    if method not in {"Manual", "Camera", "QR"}:
        raise HTTPException(status_code=400, detail="Invalid attendance method")
    subject = subject.strip()
    if not subject:
        raise HTTPException(status_code=400, detail="Select a subject")
    with engine.begin() as connection:
        validate_attendance_policy(connection, parsed_date)
        student = connection.execute(text("SELECT course, course_id FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")
        if student["course_id"] != class_id:
            raise HTTPException(status_code=403, detail="Student does not belong to the selected class")
        require_teacher_course(request, connection, student["course"], subject)
        old = connection.execute(text("SELECT id, status FROM attendance WHERE student_id = :student_id AND attendance_date = :date AND subject = :subject"), {"student_id": student_id, "date": parsed_date, "subject": subject}).mappings().first()
        row = connection.execute(text("INSERT INTO attendance (student_id, attendance_date, subject, status, method) VALUES (:student_id, :date, :subject, :status, :method) ON CONFLICT (student_id, attendance_date, subject) DO UPDATE SET status = EXCLUDED.status, method = EXCLUDED.method RETURNING id"), {"student_id": student_id, "date": parsed_date, "subject": subject, "status": status_value, "method": method}).mappings().first()
        connection.execute(text("INSERT INTO attendance_audit_logs (attendance_id, student_id, attendance_date, old_status, new_status, method, changed_by_role, changed_by_id) VALUES (:attendance_id, :student_id, :date, :old_status, :new_status, :method, :role, :changed_by)"), {"attendance_id": row["id"], "student_id": student_id, "date": parsed_date, "old_status": old["status"] if old else None, "new_status": status_value, "method": method, "role": request.session["role"], "changed_by": request.session["account_id"]})
    return {"status": status_value, "student_id": student_id, "subject": subject, "attendance_date": parsed_date.isoformat()}


@api.post("/api/teacher/results")
def add_result(request: Request, student_id: int = Form(...), class_id: int = Form(...), subject: str = Form(...), exam_name: str = Form(...), marks: float = Form(...), total_marks: float = Form(...)):
    require_role(request, "teacher", "admin")
    if marks < 0 or total_marks <= 0 or marks > total_marks:
        raise HTTPException(status_code=400, detail="Marks must be within the total marks")
    with engine.begin() as connection:
        student = connection.execute(text("SELECT course, course_id FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")
        if student["course_id"] != class_id:
            raise HTTPException(status_code=403, detail="Student does not belong to the selected class")
        subject = subject.strip()
        require_teacher_course(request, connection, student["course"], subject)
        row = connection.execute(text("INSERT INTO results (student_id, subject, exam_name, marks, total_marks, grade) VALUES (:student_id, :subject, :exam_name, :marks, :total_marks, :grade) ON CONFLICT (student_id, subject, exam_name) DO UPDATE SET marks = EXCLUDED.marks, total_marks = EXCLUDED.total_marks, grade = EXCLUDED.grade RETURNING id, student_id, subject, exam_name, marks, total_marks, grade"), {"student_id": student_id, "subject": subject, "exam_name": exam_name, "marks": marks, "total_marks": total_marks, "grade": grade_for(marks, total_marks)}).mappings().first()
    return dict(row)


@api.post("/api/teacher/leave/{leave_id}/review")
def review_leave(request: Request, leave_id: int, status_value: str = Form(...), class_id: int = Form(...)):
    require_role(request, "teacher", "admin")
    if status_value not in {"Pending", "Approved", "Rejected"}:
        raise HTTPException(status_code=400, detail="Invalid leave status")
    with engine.begin() as connection:
        query = "UPDATE leave_requests SET status = :status, reviewed_at = CURRENT_TIMESTAMP WHERE id = :leave_id"
        params = {"status": status_value, "leave_id": leave_id}
        if request.session["role"] == "teacher":
            query += f" AND teacher_id = :teacher_id AND EXISTS (SELECT 1 FROM students s WHERE s.id = leave_requests.student_id AND s.course_id = :class_id AND {teacher_student_scope()})"
            params["teacher_id"] = request.session["account_id"]
            params["class_id"] = class_id
        row = connection.execute(text(query + " RETURNING id, status"), params).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Leave request not found")
    return dict(row)


class FaceAttendancePayload(BaseModel):
    student_id: int
    class_id: int
    subject: str
    attendance_date: str


class FaceEnrollmentPayload(BaseModel):
    descriptor: list[float]
    class_id: int
    face_url: str = ""


@api.post("/api/teacher/students/{student_id}/face")
def enroll_student_face(request: Request, student_id: int, payload: FaceEnrollmentPayload):
    require_role(request, "teacher", "admin")
    if len(payload.descriptor) != 128:
        raise HTTPException(status_code=400, detail="A 128-value face descriptor is required")
    with engine.begin() as connection:
        student = connection.execute(text("SELECT course, course_id FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")
        if student["course_id"] != payload.class_id:
            raise HTTPException(status_code=403, detail="Student does not belong to the selected class")
        require_teacher_course(request, connection, student["course"])
        connection.execute(text("INSERT INTO student_faces (student_id, descriptor, face_url) VALUES (:student_id, CAST(:descriptor AS jsonb), :face_url) ON CONFLICT (student_id) DO UPDATE SET descriptor = EXCLUDED.descriptor, face_url = EXCLUDED.face_url, updated_at = CURRENT_TIMESTAMP"), {"student_id": student_id, "descriptor": json.dumps(payload.descriptor), "face_url": payload.face_url.strip()})
    return {"student_id": student_id, "enrolled": True, "face_url": payload.face_url.strip()}


@api.get("/api/teacher/face-students")
def face_students(request: Request, class_id: int | None = None, subject: str = ""):
    require_role(request, "teacher", "admin")
    with engine.connect() as connection:
        if request.session["role"] == "admin":
            rows = connection.execute(text("SELECT s.id, s.name, s.roll_number, f.descriptor, f.face_url FROM students s JOIN student_faces f ON f.student_id = s.id ORDER BY s.name")).mappings().all()
        elif class_id is not None:
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=404, detail="Class not found")
            if not subject.strip():
                raise HTTPException(status_code=400, detail="Select a subject")
            require_teacher_course(request, connection, class_row["name"], subject.strip())
            rows = connection.execute(text("SELECT s.id, s.name, s.roll_number, f.descriptor, f.face_url FROM students s JOIN student_faces f ON f.student_id = s.id WHERE s.course_id = :class_id ORDER BY s.name"), {"class_id": class_id}).mappings().all()
        else:
            if class_id is None:
                raise HTTPException(status_code=400, detail="Select a class")
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=404, detail="Class not found")
            require_teacher_course(request, connection, class_row["name"])
            rows = connection.execute(text("SELECT s.id, s.name, s.roll_number, f.descriptor, f.face_url FROM students s JOIN student_faces f ON f.student_id = s.id WHERE s.course_id = :class_id ORDER BY s.name"), {"class_id": class_id}).mappings().all()
    return {"students": [dict(row) for row in rows]}


@api.post("/api/teacher/face-attendance")
def save_face_attendance(request: Request, payload: FaceAttendancePayload):
    require_role(request, "teacher", "admin")
    parsed_date = parse_date(payload.attendance_date, "Attendance date")
    with engine.begin() as connection:
        validate_attendance_policy(connection, parsed_date)
        enrolled = connection.execute(text("SELECT s.course, s.course_id FROM student_faces f JOIN students s ON s.id = f.student_id WHERE f.student_id = :student_id"), {"student_id": payload.student_id}).mappings().first()
        if not enrolled:
            raise HTTPException(status_code=404, detail="Student face is not enrolled")
        if enrolled["course_id"] != payload.class_id:
            raise HTTPException(status_code=403, detail="Student does not belong to the selected class")
        subject = payload.subject.strip()
        require_teacher_course(request, connection, enrolled["course"], subject)
        old = connection.execute(text("SELECT id, status FROM attendance WHERE student_id = :student_id AND attendance_date = :date AND subject = :subject"), {"student_id": payload.student_id, "date": parsed_date, "subject": subject}).mappings().first()
        row = connection.execute(text("INSERT INTO attendance (student_id, attendance_date, subject, status, method) VALUES (:student_id, :date, :subject, 'Present', 'Camera') ON CONFLICT (student_id, attendance_date, subject) DO UPDATE SET status = 'Present', method = 'Camera' RETURNING id"), {"student_id": payload.student_id, "date": parsed_date, "subject": subject}).mappings().first()
        connection.execute(text("INSERT INTO attendance_audit_logs (attendance_id, student_id, attendance_date, old_status, new_status, method, changed_by_role, changed_by_id) VALUES (:attendance_id, :student_id, :date, :old_status, 'Present', 'Camera', :role, :changed_by)"), {"attendance_id": row["id"], "student_id": payload.student_id, "date": parsed_date, "old_status": old["status"] if old else None, "role": request.session["role"], "changed_by": request.session["account_id"]})
    return {"student_id": payload.student_id, "status": "Present", "method": "Camera", "subject": payload.subject, "attendance_date": parsed_date.isoformat()}


@api.put("/api/teacher/students/{student_id}")
def update_student(request: Request, student_id: int, class_id: int = Form(...), name: str = Form(...), roll_number: str = Form(...), registration_number: str = Form(...), course: str = Form(...), email: str = Form("")):
    require_role(request, "teacher", "admin")
    with engine.begin() as connection:
        student = connection.execute(text("SELECT course, course_id FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")
        if student["course_id"] != class_id:
            raise HTTPException(status_code=403, detail="Student does not belong to the selected class")
        course_row = connection.execute(text("SELECT id FROM course_catalog WHERE id = :class_id AND name = :course AND type = 'Class'"), {"class_id": class_id, "course": course.strip()}).mappings().first()
        if not course_row:
            raise HTTPException(status_code=400, detail="Select the student's current class")
        require_teacher_course(request, connection, student["course"])
        require_teacher_course(request, connection, course.strip())
        row = connection.execute(text("UPDATE students SET name=:name, roll_number=:roll, registration_number=:registration, course=:course, course_id=:course_id, email=:email WHERE id=:id RETURNING id, name, roll_number, registration_number, course, email"), {"id": student_id, "name": name.strip(), "roll": roll_number.strip(), "registration": registration_number.strip(), "course": course.strip(), "course_id": course_row["id"], "email": email.strip()}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Student not found")
    return dict(row)


@api.delete("/api/teacher/students/{student_id}")
def delete_student(request: Request, student_id: int, class_id: int):
    require_role(request, "teacher", "admin")
    with engine.begin() as connection:
        student = connection.execute(text("SELECT course, course_id FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student not found")
        if student["course_id"] != class_id:
            raise HTTPException(status_code=403, detail="Student does not belong to the selected class")
        require_teacher_course(request, connection, student["course"])
        row = connection.execute(text("DELETE FROM students WHERE id=:id RETURNING id, name"), {"id": student_id}).mappings().first()
    if not row:
        raise HTTPException(status_code=404, detail="Student not found")
    return {"deleted": True, **dict(row)}


@api.get("/api/teacher/qr")
def create_qr_session(request: Request, attendance_date: str = "", course: str = "", subject: str = ""):
    require_role(request, "teacher", "admin")
    attendance_date = parse_date(attendance_date or datetime.now().date().isoformat(), "Attendance date")
    course = course.strip()
    subject = subject.strip()
    if request.session["role"] == "teacher" and (not course or not subject):
        raise HTTPException(status_code=400, detail="Select an assigned class and subject before generating a QR code")
    course_id = None
    with engine.connect() as connection:
        validate_attendance_policy(connection, attendance_date)
    if course:
        with engine.connect() as connection:
            course_row = connection.execute(text("SELECT id FROM course_catalog WHERE name = :course AND type = 'Class'"), {"course": course}).mappings().first()
            if not course_row:
                raise HTTPException(status_code=400, detail="Select an existing class")
            course_id = course_row["id"]
            subject_row = connection.execute(text("SELECT 1 FROM course_catalog WHERE name = :subject AND type = 'Subject'"), {"subject": subject}).first()
            if not subject_row:
                raise HTTPException(status_code=400, detail="Select an existing subject")
            require_teacher_course(request, connection, course, subject)
    token_id = secrets.token_hex(16)
    token = qr_serializer.dumps({"token_id": token_id, "attendance_date": attendance_date.isoformat(), "course": course, "subject": subject})
    scan_url = f"{os.getenv('ATTENDANCE_BASE_URL', 'http://127.0.0.1:8000').rstrip('/')}/qr/scan?token={token}"
    image = qrcode.make(scan_url)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    with engine.begin() as connection:
        connection.execute(text("INSERT INTO qr_sessions (token_id, attendance_date, course, course_id, subject, created_at) VALUES (:token_id, :date, :course, :course_id, :subject, CURRENT_TIMESTAMP)"), {"token_id": token_id, "date": attendance_date, "course": course, "course_id": course_id, "subject": subject})
    return {"token": token, "scan_url": scan_url, "qr_data": base64.b64encode(buffer.getvalue()).decode("ascii"), "attendance_date": attendance_date.isoformat(), "course": course, "subject": subject}


@api.post("/api/qr/verify")
def verify_qr_attendance(token: str = Form(...), roll_number: str = Form(...), registration_number: str = Form(...)):
    try:
        payload = qr_serializer.loads(token, max_age=int(os.getenv("QR_MAX_AGE_SECONDS", "86400")))
    except (BadSignature, TypeError):
        raise HTTPException(status_code=400, detail="QR code is invalid or expired")
    with engine.begin() as connection:
        qr_session = connection.execute(text("SELECT id, used, attendance_date, course, course_id, subject FROM qr_sessions WHERE token_id=:token_id FOR UPDATE"), {"token_id": payload.get("token_id")}).mappings().first()
        if not qr_session or qr_session["used"]:
            raise HTTPException(status_code=400, detail="QR code is already used or invalid")
        student = connection.execute(text("SELECT id, name, roll_number FROM students WHERE roll_number=:roll AND registration_number=:registration"), {"roll": roll_number.strip(), "registration": registration_number.strip()}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student details did not match")
        if qr_session["course_id"] and connection.execute(text("SELECT 1 FROM students WHERE id=:id AND course_id=:course_id"), {"id": student["id"], "course_id": qr_session["course_id"]}).first() is None:
            raise HTTPException(status_code=403, detail="Student is not enrolled in this class")
        old = connection.execute(text("SELECT id, status FROM attendance WHERE student_id = :student_id AND attendance_date = :date AND subject = :subject"), {"student_id": student["id"], "date": qr_session["attendance_date"], "subject": qr_session["subject"]}).mappings().first()
        attendance = connection.execute(text("INSERT INTO attendance (student_id, attendance_date, subject, status, method) VALUES (:student_id, :date, :subject, 'Present', 'QR') ON CONFLICT (student_id, attendance_date, subject) DO UPDATE SET status='Present', method='QR' RETURNING id"), {"student_id": student["id"], "date": qr_session["attendance_date"], "subject": qr_session["subject"]}).mappings().first()
        connection.execute(text("INSERT INTO attendance_audit_logs (attendance_id, student_id, attendance_date, old_status, new_status, method, changed_by_role, changed_by_id) VALUES (:attendance_id, :student_id, :date, :old_status, 'Present', 'QR', 'student', :student_id)"), {"attendance_id": attendance["id"], "student_id": student["id"], "date": qr_session["attendance_date"], "old_status": old["status"] if old else None})
        connection.execute(text("UPDATE qr_sessions SET used=TRUE, used_by_student_id=:student_id, used_at=CURRENT_TIMESTAMP WHERE id=:id"), {"student_id": student["id"], "id": qr_session["id"]})
    return {"student": dict(student), "status": "Present", "attendance_date": str(qr_session["attendance_date"]), "method": "QR", "subject": qr_session["subject"]}


@api.get("/api/teacher/attendance/audit")
def attendance_audit(request: Request, student_id: int | None = None, class_id: int | None = None):
    require_role(request, "teacher", "admin")
    with engine.connect() as connection:
        query = "SELECT a.id, a.student_id, s.name student_name, s.course, marked.subject, a.attendance_date, a.old_status, a.new_status, a.method, a.changed_by_role, a.changed_by_id, a.changed_at FROM attendance_audit_logs a JOIN students s ON s.id = a.student_id LEFT JOIN attendance marked ON marked.id = a.attendance_id"
        params = {}
        filters = []
        if student_id is not None:
            filters.append("a.student_id = :student_id")
            params["student_id"] = student_id
        if request.session["role"] == "teacher":
            if class_id is None:
                raise HTTPException(status_code=400, detail="Select a class")
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=404, detail="Class not found")
            require_teacher_course(request, connection, class_row["name"])
            filters.append("(s.course_id = :class_id OR (s.course_id IS NULL AND s.course = :class_name))")
            params["class_id"] = class_id
            params["class_name"] = class_row["name"]
        if filters:
            query += " WHERE " + " AND ".join(filters)
        query += " ORDER BY a.changed_at DESC LIMIT 500"
        rows = connection.execute(text(query), params).mappings().all()
    return {"items": [dict(row) for row in rows]}


@api.api_route("/qr/scan", methods=["GET", "POST"], response_class=HTMLResponse, operation_id="qr_scan_page")
def qr_scan_page(token: str = "", roll_number: str = Form(""), registration_number: str = Form("")):
    if not token:
        return HTMLResponse("<h1>Invalid QR code</h1><p>Please scan a fresh attendance QR.</p>", status_code=400)
    if not roll_number or not registration_number:
        return HTMLResponse(f"""
            <title>Verify attendance</title><h1>Verify your attendance</h1>
            <p>Enter your registered roll number and registration number.</p>
            <form method="post"><input type="hidden" name="token" value="{token}">
            <label>Roll number <input name="roll_number" required></label><br>
            <label>Registration number <input name="registration_number" required></label><br>
            <button type="submit">Verify and mark Present</button></form>
        """)
    try:
        result = verify_qr_attendance(token=token, roll_number=roll_number, registration_number=registration_number)
        return HTMLResponse(f"<h1>Attendance marked</h1><p>{result['student']['name']} is Present for {result['attendance_date']}.</p>")
    except HTTPException as error:
        return HTMLResponse(f"<h1>Attendance not marked</h1><p>{error.detail}</p>", status_code=error.status_code)


@api.get("/api/student/overview")
def student_overview(request: Request):
    require_role(request, "student")
    student_id = request.session["account_id"]
    with engine.connect() as connection:
        student = connection.execute(text("SELECT id, name, roll_number, registration_number, course, email FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        attendance = connection.execute(text("SELECT id, attendance_date, subject, status, method FROM attendance WHERE student_id = :id ORDER BY attendance_date DESC, subject"), {"id": student_id}).mappings().all()
        results = connection.execute(text("SELECT subject, exam_name, marks, total_marks, grade FROM results WHERE student_id = :id ORDER BY id DESC"), {"id": student_id}).mappings().all()
        teachers = connection.execute(text("SELECT id, name, username FROM teachers ORDER BY name")).mappings().all()
        leave_requests = connection.execute(text("SELECT lr.id, lr.date_from, lr.date_to, lr.reason, lr.status, lr.created_at, t.name teacher_name FROM leave_requests lr LEFT JOIN teachers t ON t.id = lr.teacher_id WHERE lr.student_id = :id ORDER BY lr.created_at DESC"), {"id": student_id}).mappings().all()
    return {"student": dict(student), "attendance": [dict(row) for row in attendance], "results": [dict(row) for row in results], "teachers": [dict(row) for row in teachers], "leave_requests": [dict(row) for row in leave_requests]}


@api.get("/api/student/attendance-export")
def export_student_attendance(request: Request, start_date: str, end_date: str, export_format: str):
    require_role(request, "student")
    start = parse_date(start_date, "Start date")
    end = parse_date(end_date, "End date")
    if start > end:
        raise HTTPException(status_code=400, detail="Start date must be before end date")
    if export_format not in {"pdf", "xlsx"}:
        raise HTTPException(status_code=400, detail="Export format must be pdf or xlsx")

    student_id = request.session["account_id"]
    with engine.connect() as connection:
        student = connection.execute(text("SELECT name, roll_number, course FROM students WHERE id = :id"), {"id": student_id}).mappings().first()
        if not student:
            raise HTTPException(status_code=404, detail="Student profile was not found")
        rows = connection.execute(text("SELECT attendance_date, subject, status, method FROM attendance WHERE student_id = :id AND attendance_date BETWEEN :start AND :end ORDER BY attendance_date, subject"), {"id": student_id, "start": start, "end": end}).mappings().all()

    filename = f"attendance-{start.isoformat()}-to-{end.isoformat()}"
    output = io.BytesIO()
    if export_format == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Attendance"
        sheet.append(["Student", student["name"]])
        sheet.append(["Roll number", student["roll_number"]])
        sheet.append(["Class", student["course"]])
        sheet.append([])
        sheet.append(["Date", "Subject", "Status", "Method"])
        for cell in sheet[5]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="145B58")
        for row in rows:
            sheet.append([row["attendance_date"], row["subject"], row["status"], row["method"]])
        sheet.freeze_panes = "A6"
        sheet.auto_filter.ref = f"A5:D{max(sheet.max_row, 5)}"
        for column, width in {"A": 16, "B": 28, "C": 16, "D": 16}.items():
            sheet.column_dimensions[column].width = width
        workbook.save(output)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        extension = "xlsx"
    else:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        document = SimpleDocTemplate(output, pagesize=landscape(A4))
        styles = getSampleStyleSheet()
        table_data = [["Date", "Subject", "Status", "Method"]]
        table_data.extend([[row["attendance_date"].isoformat(), row["subject"], row["status"], row["method"]] for row in rows])
        table = Table(table_data, repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#145B58")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#D5DEDC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6F5")]),
            ("PADDING", (0, 0), (-1, -1), 8),
        ]))
        document.build([
            Paragraph("Student Attendance", styles["Title"]),
            Paragraph(f"{student['name']} | {student['roll_number']} | {student['course']}", styles["Normal"]),
            Paragraph(f"Date range: {start.isoformat()} to {end.isoformat()}", styles["Normal"]),
            Spacer(1, 16),
            table,
        ])
        media_type = "application/pdf"
        extension = "pdf"

    output.seek(0)
    return Response(output.getvalue(), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}.{extension}"'})


@api.post("/api/student/leave")
def submit_leave(request: Request, teacher_id: int = Form(...), date_from: str = Form(...), date_to: str = Form(...), reason: str = Form(...)):
    require_role(request, "student")
    try:
        start_date = datetime.strptime(date_from, "%Y-%m-%d").date()
        end_date = datetime.strptime(date_to, "%Y-%m-%d").date()
    except ValueError as error:
        raise HTTPException(status_code=400, detail="Please select valid leave dates") from error
    if start_date > end_date or not reason.strip():
        raise HTTPException(status_code=400, detail="Select a valid date range and write a reason")
    with engine.begin() as connection:
        teacher = connection.execute(text("SELECT id FROM teachers WHERE id = :teacher_id"), {"teacher_id": teacher_id}).first()
        if not teacher:
            raise HTTPException(status_code=404, detail="Selected teacher was not found")
        row = connection.execute(text("INSERT INTO leave_requests (student_id, teacher_id, date_from, date_to, reason) VALUES (:student_id, :teacher_id, :date_from, :date_to, :reason) RETURNING id, status"), {"student_id": request.session["account_id"], "teacher_id": teacher_id, "date_from": start_date, "date_to": end_date, "reason": reason.strip()}).mappings().first()
        connection.execute(text("INSERT INTO notifications (recipient_role, recipient_id, title, body) VALUES ('teacher', :teacher_id, 'New leave request', :body)"), {"teacher_id": teacher_id, "body": f"A student submitted leave from {start_date} to {end_date}."})
    return dict(row)


@api.get("/api/reports/attendance.csv")
def attendance_report(request: Request, start_date: str = "", end_date: str = "", class_id: int | None = None):
    require_role(request, "admin", "teacher")
    start = parse_date(start_date or datetime.now().date().isoformat(), "Start date")
    end = parse_date(end_date or start.isoformat(), "End date")
    if start > end:
        raise HTTPException(status_code=400, detail="Start date must be before end date")
    with engine.connect() as connection:
        params = {"start": start, "end": end}
        query = "SELECT s.name, s.roll_number, s.course, a.subject, a.attendance_date, a.status, a.method FROM attendance a JOIN students s ON s.id = a.student_id WHERE a.attendance_date BETWEEN :start AND :end"
        if request.session["role"] == "teacher":
            if class_id is None:
                raise HTTPException(status_code=400, detail="Select a class")
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=404, detail="Class not found")
            require_teacher_course(request, connection, class_row["name"])
            query += " AND (s.course_id = :class_id OR (s.course_id IS NULL AND s.course = :class_name))"
            params["class_id"] = class_id
            params["class_name"] = class_row["name"]
        rows = connection.execute(text(query + " ORDER BY a.attendance_date, s.name"), params).mappings().all()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Student", "Roll number", "Class", "Subject", "Date", "Status", "Method"])
    writer.writerows([row.values() for row in rows])
    return StreamingResponse(iter([output.getvalue()]), media_type="text/csv", headers={"Content-Disposition": "attachment; filename=attendance-report.csv"})


@api.get("/api/notifications")
def notifications(request: Request):
    require_session(request)
    with engine.connect() as connection:
        rows = connection.execute(text("SELECT id, title, body, read_at, created_at FROM notifications WHERE recipient_role = :role AND recipient_id = :id ORDER BY created_at DESC LIMIT 100"), {"role": request.session["role"], "id": request.session["account_id"]}).mappings().all()
    return {"notifications": [dict(row) for row in rows]}


@api.get("/api/dashboard")
def dashboard_data(request: Request, class_id: int | None = None):
    require_role(request, "admin", "teacher")
    with engine.connect() as connection:
        if request.session["role"] == "teacher":
            if class_id is None:
                raise HTTPException(status_code=400, detail="Select a class")
            class_row = connection.execute(text("SELECT id, name FROM course_catalog WHERE id = :id AND type = 'Class'"), {"id": class_id}).mappings().first()
            if not class_row:
                raise HTTPException(status_code=404, detail="Class not found")
            require_teacher_course(request, connection, class_row["name"])
            params = {"class_id": class_id, "class_name": class_row["name"], "teacher_id": request.session["account_id"]}
            student_scope = "(s.course_id = :class_id OR (s.course_id IS NULL AND s.course = :class_name))"
            course_count = "(SELECT COUNT(*) FROM (SELECT course_id FROM teacher_class_courses WHERE teacher_id = :teacher_id UNION SELECT course_id FROM teacher_courses WHERE teacher_id = :teacher_id) assigned_courses)"
        else:
            params = {}
            student_scope = "TRUE"
            course_count = "(SELECT COUNT(*) FROM course_catalog)"
        totals = connection.execute(text(f"""
            SELECT
                (SELECT COUNT(*) FROM students s WHERE {student_scope}) AS students,
                {course_count} AS courses,
                (SELECT COUNT(*) FROM attendance a JOIN students s ON s.id = a.student_id WHERE a.attendance_date = CURRENT_DATE AND a.status = 'Present' AND {student_scope}) AS present,
                (SELECT COUNT(*) FROM attendance a JOIN students s ON s.id = a.student_id WHERE a.attendance_date = CURRENT_DATE AND a.status = 'Absent' AND {student_scope}) AS absent,
                (SELECT COUNT(*) FROM attendance a JOIN students s ON s.id = a.student_id WHERE a.attendance_date = CURRENT_DATE AND a.status = 'Late' AND {student_scope}) AS late,
                (SELECT COUNT(*) FROM leave_requests lr JOIN students s ON s.id = lr.student_id WHERE lr.status = 'Pending' AND {student_scope}) AS pending_leave_requests
        """), params).mappings().first()
        course_rows = connection.execute(text(f"""
            SELECT s.course,
                   COUNT(DISTINCT s.id) AS total_students,
                   COUNT(*) FILTER (WHERE a.status = 'Present') AS present,
                   COUNT(*) FILTER (WHERE a.status = 'Absent') AS absent,
                   COUNT(*) FILTER (WHERE a.status = 'Late') AS late
            FROM students s
            LEFT JOIN attendance a ON a.student_id = s.id AND a.attendance_date = CURRENT_DATE
            WHERE {student_scope}
            GROUP BY s.course ORDER BY s.course
        """), params).mappings().all()
        recent = connection.execute(text(f"""
            SELECT s.name, s.roll_number, a.attendance_date, a.status
            FROM attendance a JOIN students s ON s.id = a.student_id
            WHERE a.attendance_date = CURRENT_DATE AND {student_scope} ORDER BY a.id DESC LIMIT 8
        """), params).mappings().all()
    courses = []
    for row in course_rows:
        marked = row["present"] + row["absent"] + row["late"]
        percentage = round(row["present"] / marked * 100, 1) if marked else 0
        courses.append({**dict(row), "percentage": percentage})
    marked = totals["present"] + totals["absent"] + totals["late"]
    percentage = round(totals["present"] / marked * 100, 1) if marked else 0
    return {"stats": {**dict(totals), "average_attendance": percentage, "students_below_75": 0}, "courses": courses, "recent": [dict(row) for row in recent]}

app = api
